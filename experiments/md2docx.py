"""Minimal Markdown -> DOCX (headings, bullets, numbered lists, tables, **bold**, `code`)."""
import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT

src, dst = sys.argv[1], sys.argv[2]
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.6))
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10)

def runs(par, text, size=None):
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok: continue
        if tok.startswith("**"): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"): r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
        else: r = par.add_run(tok)
        if size: r.font.size = Pt(size)

lines = open(src, encoding="utf8").read().split("\n")
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); t.style = "Light Grid Accent 1"
        for ri, r in enumerate(rows):
            for ci, c in enumerate(r):
                cell = t.cell(ri, ci); cell.text = ""
                runs(cell.paragraphs[0], c, 8.5)
                if ri == 0:
                    for rr in cell.paragraphs[0].runs: rr.bold = True
        doc.add_paragraph(); continue
    m = re.match(r"(#+) (.*)", ln)
    if m: doc.add_heading(m.group(2).replace("**", ""), min(len(m.group(1)) - 1, 3) if len(m.group(1)) > 1 else 0)
    elif ln.strip() == "---": pass
    elif re.match(r"\s*- ", ln): runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*- ", "", ln))
    elif re.match(r"\d+\. ", ln): runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", ln))
    elif ln.strip(): runs(doc.add_paragraph(), ln)
    i += 1
doc.save(dst)
